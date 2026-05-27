import logging
from typing import List

from extractors.base import BaseExtractor, ExtractedUrl
from utils.url_utils import extract_plain_urls, get_context

logger = logging.getLogger(__name__)


class XlsxExtractor(BaseExtractor):
    def extract(self):
        # type: () -> List[ExtractedUrl]
        try:
            from openpyxl import load_workbook
        except ImportError:
            logger.error("openpyxl not installed")
            return []

        try:
            wb = load_workbook(self.file_path, read_only=True, data_only=True)
        except Exception as e:
            logger.error("Cannot open %s: %s", self.file_path, e)
            return []

        results = []  # type: List[ExtractedUrl]

        for sheet_name in wb.sheetnames:
            try:
                ws = wb[sheet_name]
                for row in ws.iter_rows():
                    for cell in row:
                        coord = cell.coordinate
                        location = "Sheet '{}', {}".format(sheet_name, coord)

                        try:
                            hl = cell.hyperlink
                            if hl:
                                url = getattr(hl, "target", None) or getattr(hl, "url", None)
                                if url:
                                    results.append(ExtractedUrl(
                                        url=url,
                                        source_file=self.source_name,
                                        location=location,
                                        context=str(cell.value or "")[:100],
                                        link_type="embedded",
                                    ))
                        except Exception:
                            pass

                        val = cell.value
                        if val and isinstance(val, str) and val.strip():
                            for url, pos in extract_plain_urls(val):
                                results.append(ExtractedUrl(
                                    url=url,
                                    source_file=self.source_name,
                                    location=location,
                                    context=get_context(val, pos),
                                    link_type="plain-text",
                                ))
            except Exception as e:
                logger.warning("Error reading sheet '%s' in %s: %s", sheet_name, self.source_name, e)

        return results
